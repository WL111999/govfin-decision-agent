"""表格解析：CSV / TSV / 类 XLSX 的管道分隔文本。

含表头推断与类型嗅探——把"1,234.56"这种千分位数字还原成 float，
否则下游建图时"注册资本"会变成字符串，SHACL 的 datatype 约束会全部报错。
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from govfin.errors import ParseError
from govfin.graph.schema import MODALITY_TABLE
from govfin.ingest.multimodal import ParsedDocument

_NUMBER = re.compile(r"^-?\d{1,3}(?:,\d{3})*(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")
_INTEGER = re.compile(r"^-?\d+$")
_PERCENT = re.compile(r"^-?\d+(?:\.\d+)?%$")
_DATE = re.compile(r"^\d{4}[-/年]\d{1,2}(?:[-/月]\d{1,2}日?)?$")
_HEADER_CHARS = ("名称", "编号", "金额", "日期", "类型", "状态", "代码", "月份", "人数", "科目", "项目", "单位")


class TableParser:
    name = "table"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in (".csv", ".tsv", ".xlsx", ".xls")

    def parse(self, path: str | Path, *, source: str | None = None, **_) -> ParsedDocument:
        p = Path(path)
        try:
            data = p.read_bytes()
        except OSError as exc:
            raise ParseError(f"无法读取 '{p}': {exc}") from exc
        suffix = p.suffix.lower()
        if suffix in (".xlsx", ".xls"):
            return self._parse_excel_bytes(data, source=source or p.name)
        return self.parse_bytes(data, source=source or p.name, delimiter="\t" if suffix == ".tsv" else None)

    def parse_bytes(self, data: bytes, *, source: str, delimiter: str | None = None, **_) -> ParsedDocument:
        if not data.strip():
            raise ParseError(f"'{source}' 为空文件")

        text = _decode(data)
        if delimiter is None:
            delimiter = _sniff_delimiter(text)

        doc = ParsedDocument(source=source, modality=MODALITY_TABLE, metadata={"delimiter": delimiter})
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error as exc:
            raise ParseError(f"'{source}' 表格解析失败: {exc}") from exc

        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            raise ParseError(f"'{source}' 没有有效行")

        header, body = _split_header(rows)
        doc.metadata["header"] = header
        doc.metadata["row_count"] = len(body)

        typed_rows = [[_coerce(cell) for cell in row] for row in body]
        rendered_header = "|".join(header) if header else ""
        for index, row in enumerate(typed_rows):
            rendered = "|".join(_render_cell(c) for c in row)
            doc.add(
                "table",
                rendered,
                locator=f"row:{index}",
                header=header,
                row_index=index,
                values=row,
                header_line=rendered_header,
            )
        return doc

    def _parse_excel_bytes(self, data: bytes, *, source: str) -> ParsedDocument:
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise ParseError(
                "缺少 openpyxl，无法解析 Excel。请 `pip install openpyxl`，"
                "或把文件另存为 CSV 后重试。"
            ) from exc

        doc = ParsedDocument(source=source, modality=MODALITY_TABLE, metadata={"format": "xlsx"})
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Excel '{source}' 打开失败: {exc}") from exc

        for sheet in wb.worksheets:
            rows = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if not rows:
                continue
            header, body = _split_header(rows)
            for index, row in enumerate(body):
                doc.add(
                    "table",
                    "|".join(row),
                    locator=f"sheet:{sheet.title}/row:{index}",
                    sheet=sheet.title,
                    header=header,
                    row_index=index,
                    values=[_coerce(c) for c in row],
                )
        wb.close()
        return doc


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("表格文件编码无法识别")


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:20])
    candidates = {",": sample.count(","), "\t": sample.count("\t"), ";": sample.count(";"), "|": sample.count("|")}
    best = max(candidates, key=lambda k: candidates[k])
    return best if candidates[best] > 0 else ","


def _split_header(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    if not rows:
        return [], []
    first = rows[0]
    is_header = any(any(marker in cell for marker in _HEADER_CHARS) for cell in first)
    # 表头行不应含纯数字单元格
    if is_header and not any(_looks_numeric(c) for c in first if c.strip()):
        return [c.strip() for c in first], rows[1:]
    return [f"列{i + 1}" for i in range(len(first))], rows


def _looks_numeric(cell: str) -> bool:
    return bool(_NUMBER.match(cell.strip()) or _PERCENT.match(cell.strip()))


def _coerce(cell: str):
    s = (cell or "").strip()
    if not s:
        return ""
    if _PERCENT.match(s):
        return round(float(s[:-1]) / 100.0, 6)
    if _NUMBER.match(s):
        cleaned = s.replace(",", "")
        if _INTEGER.match(cleaned):
            try:
                return int(cleaned)
            except ValueError:
                pass
        try:
            return float(cleaned)
        except ValueError:
            return cell
    if _DATE.match(s):
        return s.replace("/", "-").replace("年", "-").replace("月", "-").rstrip("日").strip("-")
    return cell


def _render_cell(value) -> str:
    return f"{value:.6g}" if isinstance(value, float) else str(value)
